"""Bounded dataset import/export. Application code is MIT licensed."""

from __future__ import annotations

import csv
from contextlib import closing
import datetime
import importlib
import json
import math
import os
from pathlib import Path
import re
import sqlite3
import tempfile
import zipfile

from backend.common import (
    MAX_ROWS, MAX_COLUMNS, MAX_CELLS, MAX_FILE_BYTES, MAX_CELL_CHARACTERS,
    validate_dataset, reject_json_constant, unique_json_object, sanitize_json,
)

_NUMBER = re.compile(r"^[+-]?(?:(?:0|[1-9]\d*)(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$")
_SAFE_INTEGER = 9_007_199_254_740_991


def _format_metadata(data):
    value = data.get("formatMetadata") if isinstance(data, dict) else None
    if value is None:
        return None
    if not isinstance(value, dict):
        raise ValueError("File format metadata must be an object.")
    encoded = json.dumps(value, ensure_ascii=False, allow_nan=False)
    if len(encoded.encode("utf-8")) > 8 * 1024 * 1024:
        raise ValueError("File format metadata exceeds the 8 MiB limit.")
    return json.loads(encoded)


def _module(name, package=None):
    try:
        return importlib.import_module(name)
    except ImportError as exc:
        raise ValueError(f"This format requires the optional '{package or name}' package. Install requirements-formats.txt in the Python environment.") from exc


def _input_path(path):
    if not isinstance(path, str) or not path.strip():
        raise ValueError("A file path is required.")
    p = Path(path).expanduser().resolve()
    if not p.is_file():
        raise ValueError("The selected file does not exist or is not a regular file.")
    if p.stat().st_size > MAX_FILE_BYTES:
        raise ValueError("The file exceeds the 64 MiB import limit.")
    return p


def _text_cell(value):
    if value == "":
        return None
    stripped = value.strip()
    # Preserve identifier-like leading zeros, literal NA/NaN, and large integers.
    if stripped and _NUMBER.fullmatch(stripped):
        if not any(c in stripped.lower() for c in ".e"):
            integer = int(stripped)
            return integer if abs(integer) <= _SAFE_INTEGER else value
        number = float(stripped)
        return number if math.isfinite(number) else value
    return value


def _headers(values):
    if len(values) > MAX_COLUMNS:
        raise ValueError(f"The file contains more than {MAX_COLUMNS:,} columns.")
    result = [str(v) if v is not None and str(v).strip() else f"Column {i + 1}" for i, v in enumerate(values)]
    validate_dataset({"columns": result, "rows": []})
    return result


def _bounded_append(rows, row, width):
    if len(rows) >= MAX_ROWS or (len(rows) + 1) * width > MAX_CELLS:
        raise ValueError(f"Import limit: {MAX_ROWS:,} rows and {MAX_CELLS:,} cells.")
    rows.append(row)


def _csv_import(path):
    csv.field_size_limit(MAX_CELL_CHARACTERS)
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        sample = stream.read(8192)
        stream.seek(0)
        delimiter = "\t" if path.suffix.lower() in {".tsv", ".tab"} else ","
        if delimiter == ",":
            try:
                delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
            except csv.Error:
                pass
        reader = csv.reader(stream, delimiter=delimiter, strict=True)
        try:
            columns = _headers(next(reader))
        except StopIteration as exc:
            raise ValueError("The delimited file is empty.") from exc
        rows = []
        for row in reader:
            if not row:
                row = [""] * len(columns)
            if len(row) != len(columns):
                raise ValueError(f"Delimited row {reader.line_num} has {len(row)} fields; expected {len(columns)}.")
            _bounded_append(rows, [_text_cell(v) for v in row], len(columns))
    result = validate_dataset({"columns": columns, "rows": rows, "name": path.stem})
    result["notes"] = ["Empty delimited fields are missing values. Literal NA/NaN and leading-zero identifiers remain text."]
    return result


def _check_zip(path):
    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > 10_000 or sum(i.file_size for i in infos) > 256 * 1024 * 1024:
            raise ValueError("The workbook expands beyond the safe 256 MiB archive limit.")


def _scalar(value):
    if value is None:
        return None
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, int) and abs(value) > _SAFE_INTEGER:
        return str(value)
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _xlsx_import(path, sheet):
    _check_zip(path)
    openpyxl = _module("openpyxl")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        names = workbook.sheetnames
        if sheet is not None and sheet not in names:
            raise ValueError("Worksheet not found. Available worksheets: " + ", ".join(names))
        if not names:
            raise ValueError("The workbook has no worksheets.")
        selected = sheet or names[0]
        ws = workbook[selected]
        if (ws.max_column or 0) > MAX_COLUMNS or (ws.max_row or 0) > MAX_ROWS + 1:
            raise ValueError("Worksheet dimensions exceed the import limits.")
        iterator = ws.iter_rows(values_only=True)
        header = next(iterator, None)
        if header is None or all(v is None for v in header):
            raise ValueError("The selected worksheet has no header row.")
        columns = _headers(header)
        rows = []
        for row in iterator:
            if len(row) > len(columns):
                columns = _headers([*header, *([None] * (len(row) - len(header)))])
            _bounded_append(rows, [_scalar(v) for v in row], len(columns))
        rows = [row + [None] * (len(columns) - len(row)) for row in rows]
        result = validate_dataset({"columns": columns, "rows": rows, "name": path.stem})
        result.update(sheets=names, selected_sheet=selected, notes=["Excel formulas are not evaluated; cached results are imported. Dates/times are ISO text."])
        return result
    finally:
        workbook.close()


def _sqlite_import(path, sheet):
    # The connection's context manager controls transactions; it does not close
    # the file handle, which also prevents replacing/deleting files on Windows.
    with closing(sqlite3.connect(path.as_uri() + "?mode=ro", uri=True)) as connection:
        connection.execute("PRAGMA query_only=ON")
        connection.execute("PRAGMA trusted_schema=OFF")
        names = [row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' ORDER BY name")]
        if not names:
            raise ValueError("The database has no readable user tables.")
        if sheet is None and len(names) > 1:
            return {"name": path.stem, "sheets": names, "requires_sheet": True}
        selected = sheet if sheet is not None else names[0]
        if selected not in names:
            raise ValueError("Choose an existing SQLite table: " + ", ".join(names))
        quoted = '"' + selected.replace('"', '""') + '"'
        cursor = connection.execute(f"SELECT * FROM {quoted} LIMIT ?", (MAX_ROWS + 1,))
        columns = _headers([d[0] for d in cursor.description])
        rows = []
        for row in cursor:
            _bounded_append(rows, [_scalar(v) for v in row], len(columns))
        result = validate_dataset({"columns": columns, "rows": rows, "name": path.stem})
        result.update(sheets=names, selected_sheet=selected, notes=["SQLite was opened read-only. Binary cells are decoded as UTF-8 with replacement."])
        return result


def _column_letters(index):
    result = ""
    while index >= 0:
        index, remainder = divmod(index, 26)
        result = chr(65 + remainder) + result
        index -= 1
    return result


def _spjs_import(path, sheet):
    """Import an observed project container plus a cached-value dataset view."""
    with path.open("r", encoding="utf-8-sig") as stream:
        document = json.load(stream, parse_constant=reject_json_constant, object_pairs_hook=unique_json_object)
    workbook = document.get("workbook") if isinstance(document, dict) else None
    if not isinstance(workbook, dict) or not isinstance(workbook.get("sheets"), dict):
        raise ValueError("Unsupported .spjs structure: expected workbook.sheets.")
    sheets = workbook["sheets"]
    if not sheets or any(not isinstance(value, dict) for value in sheets.values()):
        raise ValueError("The .spjs file has no valid worksheets.")
    names = sorted(sheets, key=lambda key: (sheets[key].get("index", 0) if isinstance(sheets[key].get("index", 0), int) else 0))
    active = workbook.get("activeSheetIndex", 0)
    default = next((name for name in names if sheets[name].get("index", 0) == active), names[0])
    selected = sheet or default
    if selected not in names:
        raise ValueError("Worksheet not found. Available worksheets: " + ", ".join(names))
    content = sheets[selected].get("data", {})
    table = content.get("dataTable", {}) if isinstance(content, dict) else None
    if not isinstance(table, dict):
        raise ValueError("Unsupported .spjs cell data structure.")
    values, formulas, unsupported, maximum_row, maximum_col = {}, 0, 0, -1, -1
    for row_key, row in table.items():
        if not str(row_key).isdigit() or not isinstance(row, dict):
            raise ValueError("Invalid .spjs row index or cell row.")
        r = int(row_key)
        for col_key, cell in row.items():
            if not str(col_key).isdigit():
                # SpreadJS row metadata is not cell content.
                continue
            c = int(col_key)
            if not isinstance(cell, dict) or not ("value" in cell or "formula" in cell):
                continue
            if r >= MAX_ROWS or c >= MAX_COLUMNS:
                raise ValueError("The .spjs used-cell range exceeds dataset limits.")
            maximum_row, maximum_col = max(maximum_row, r), max(maximum_col, c)
            if (maximum_row + 1) * (maximum_col + 1) > MAX_CELLS:
                raise ValueError("The .spjs used-cell range exceeds two million cells.")
            if "formula" in cell:
                formulas += 1
            value = cell.get("value")
            if isinstance(value, dict):
                if isinstance(value.get("_calcError", value.get("_error")), str):
                    value = value.get("_calcError", value.get("_error"))
                else:
                    value, unsupported = None, unsupported + 1
            elif isinstance(value, list):
                value, unsupported = None, unsupported + 1
            values[(r, c)] = _scalar(value)
    width = max(1, maximum_col + 1)
    rows = [[values.get((r, c)) for c in range(width)] for r in range(maximum_row + 1)]
    result = validate_dataset({"columns": [_column_letters(c) for c in range(width)], "rows": rows, "name": selected})
    from backend.spjs_import import translate_workbook
    snapshot, charts, notes = translate_workbook(document, path.stem, selected, _scalar)
    if unsupported:
        notes.append(f"{unsupported} unsupported structured cell values became missing values; their originals are retained in project metadata.")
    result.update(sheets=names, selected_sheet=selected, notes=notes, workbook=snapshot, importedCharts=charts)
    return result


def _frame_dataset(frame, path, sheets=None, selected=None):
    if not hasattr(frame, "columns") or not hasattr(frame, "itertuples"):
        raise ValueError("The selected object is not a rectangular data frame.")
    if len(frame) > MAX_ROWS or len(frame.columns) > MAX_COLUMNS or frame.size > MAX_CELLS:
        raise ValueError("The imported data exceeds the dataset size limits.")
    pd = _module("pandas")
    columns = _headers(list(frame.columns))
    rows = [[None if pd.isna(v) else _scalar(v) for v in row] for row in frame.itertuples(index=False, name=None)]
    result = validate_dataset({"columns": columns, "rows": rows, "name": path.stem})
    if sheets is not None:
        result.update(sheets=sheets, selected_sheet=selected)
    return result


def _optional_import(path, sheet):
    ext = path.suffix.lower()
    pd = _module("pandas")
    if ext in {".xls", ".xlt", ".ods"}:
        engine = "odf" if ext == ".ods" else "xlrd"
        _module(engine, "odfpy" if engine == "odf" else engine)
        if ext == ".ods":
            _check_zip(path)
        with pd.ExcelFile(path, engine=engine) as book:
            selected = sheet or book.sheet_names[0]
            if selected not in book.sheet_names:
                raise ValueError("Worksheet not found: " + str(selected))
            # Read header separately to detect duplicates before pandas can rename them.
            frame = pd.read_excel(book, sheet_name=selected, header=None, nrows=MAX_ROWS + 2, keep_default_na=False)
            if frame.empty:
                raise ValueError("The worksheet is empty.")
            names = _headers(frame.iloc[0].tolist())
            frame = frame.iloc[1:].copy()
            frame.columns = names
            frame = frame.replace("", None)
            return _frame_dataset(frame, path, book.sheet_names, selected)
    if ext == ".dta":
        with pd.io.stata.StataReader(path, convert_categoricals=False) as reader:
            result = _frame_dataset(reader.read(), path)
            result["formatMetadata"] = {"format": "Stata", "dataLabel": reader.data_label, "variableLabels": reader.variable_labels(), "valueLabels": sanitize_json(reader.value_labels())}
        return result
    if ext in {".sas7bdat", ".xport", ".xpt"}:
        fmt = "sas7bdat" if ext == ".sas7bdat" else "xport"
        return _frame_dataset(pd.read_sas(path, format=fmt), path)
    if ext in {".sav", ".zsav", ".por"}:
        module = _module("pyreadstat")
        read = module.read_por if ext == ".por" else module.read_sav
        frame, metadata = read(str(path), row_limit=MAX_ROWS + 1, **({"user_missing": True} if ext != ".por" else {}))
        missing = []
        for column, ranges in (metadata.missing_ranges or {}).items():
            for index, value in enumerate(frame[column]):
                if any(item["lo"] <= value <= item["hi"] for item in ranges if not pd.isna(value)):
                    missing.append({"row": index, "column": column, "value": _scalar(value)})
                    frame.at[index, column] = None
        result = _frame_dataset(frame, path)
        fields = ("column_names_to_labels", "variable_value_labels", "missing_ranges", "missing_user_values", "original_variable_types", "variable_measure", "variable_display_width", "file_label", "notes", "file_encoding")
        result["formatMetadata"] = {"format": "SPSS", **{key: sanitize_json(getattr(metadata, key)) for key in fields if getattr(metadata, key, None) is not None}, "userMissingCells": missing}
        result["notes"] = ["SPSS declared user-missing codes become missing cells; original codes, labels and declarations remain in format metadata."]
        return result
    if ext == ".dbf":
        module = _module("dbfread")
        table = module.DBF(str(path), load=False)
        columns = _headers(table.field_names)
        rows = []
        for record in table:
            _bounded_append(rows, [_scalar(record[name]) for name in table.field_names], len(columns))
        return validate_dataset({"columns": columns, "rows": rows, "name": path.stem})
    if ext in {".rdata", ".rda", ".rds"}:
        # pyreadr parses supported data objects in C; no R interpreter or eval is invoked.
        module = _module("pyreadr")
        objects = module.read_r(str(path))
        if not objects:
            raise ValueError("No supported data frames were found in the R file.")
        names = [str(key) if key is not None else "Data" for key in objects]
        if sheet is None and len(names) > 1:
            return {"name": path.stem, "sheets": names, "requires_sheet": True}
        selected = sheet or names[0]
        if selected not in names:
            raise ValueError("Choose an R data object: " + ", ".join(names))
        frame = list(objects.values())[names.index(selected)]
        return _frame_dataset(frame, path, names, selected)
    raise ValueError("Unsupported file format. Use CSV, TSV, XLSX, dataset JSON, or a supported optional statistical format.")


def _import_file(path, sheet=None):
    path = _input_path(path)
    if sheet is not None and not isinstance(sheet, str):
        raise ValueError("Worksheet/table name must be text.")
    ext = path.suffix.lower()
    if ext in {".csv", ".tsv", ".tab", ".txt"}:
        return _csv_import(path)
    if ext == ".json":
        with path.open("r", encoding="utf-8-sig") as stream:
            data = json.load(stream, parse_constant=reject_json_constant, object_pairs_hook=unique_json_object)
        result = validate_dataset(data)
        if (metadata := _format_metadata(data)) is not None:
            result["formatMetadata"] = metadata
        return result
    if ext == ".xlsx":
        return _xlsx_import(path, sheet)
    if ext == ".spjs":
        return _spjs_import(path, sheet)
    if ext in {".db", ".db3", ".sqlite", ".sqlite3"}:
        return _sqlite_import(path, sheet)
    return _optional_import(path, sheet)


def import_file(path, sheet=None):
    try:
        return _import_file(path, sheet)
    except UnicodeDecodeError as exc:
        raise ValueError("The text file is not valid UTF-8. Save it as UTF-8 CSV/TSV or use XLSX.") from exc
    except csv.Error as exc:
        raise ValueError(f"Invalid delimited file: {exc}") from exc
    except zipfile.BadZipFile as exc:
        raise ValueError("The selected workbook is not a valid ZIP-based XLSX/ODS file.") from exc
    except sqlite3.DatabaseError as exc:
        raise ValueError(f"The selected SQLite database could not be read: {exc}") from exc


def export_file(path, data):
    """Write atomically; XLSX text is always text, never an injected formula."""
    metadata = _format_metadata(data)
    data = validate_dataset(data)
    if not isinstance(path, str) or not path.strip():
        raise ValueError("An output path is required.")
    path = Path(path).expanduser().resolve()
    if not path.parent.is_dir():
        raise ValueError("The output folder does not exist.")
    ext = path.suffix.lower()
    if ext not in {".csv", ".tsv", ".xlsx", ".json"}:
        raise ValueError("Export supports CSV, TSV, XLSX and dataset JSON.")
    if path.name.lower().endswith(".statworkbench.json"):
        raise ValueError("Use Save Project for .statworkbench.json files; choose another name for dataset JSON.")
    temporary = None
    try:
        fd, temporary = tempfile.mkstemp(prefix=".statworkbench-", suffix=ext, dir=path.parent)
        os.close(fd)
        temp = Path(temporary)
        notes = []
        if metadata is not None and ext != ".json":
            notes.append("This dataset export contains cell values; use dataset JSON or Save Project to retain imported statistical-file metadata.")
        if ext in {".csv", ".tsv"}:
            with temp.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.writer(stream, delimiter="\t" if ext == ".tsv" else ",")
                writer.writerow(data["columns"])
                writer.writerows(data["rows"])
            notes.append("Delimited export cannot distinguish empty strings from missing cells; dataset JSON preserves this distinction.")
            if any(isinstance(v, str) and v.lstrip().startswith(("=", "+", "-", "@")) for row in [data["columns"], *data["rows"]] for v in row):
                notes.append("Some text begins with spreadsheet formula characters. Import CSV/TSV as text, or use XLSX export which explicitly stores text cells.")
        elif ext == ".json":
            if metadata is not None:
                data["formatMetadata"] = metadata
            with temp.open("w", encoding="utf-8") as stream:
                json.dump(data, stream, ensure_ascii=False, allow_nan=False, indent=2)
                stream.write("\n")
        else:
            module = _module("openpyxl")
            workbook = module.Workbook(write_only=True)
            ws = workbook.create_sheet("Data")
            from openpyxl.cell import WriteOnlyCell
            for row in [data["columns"], *data["rows"]]:
                cells = []
                for value in row:
                    if isinstance(value, str) and len(value) > 32_767:
                        raise ValueError("Excel cells cannot store more than 32,767 characters; use JSON export.")
                    cell = WriteOnlyCell(ws, value=value)
                    if isinstance(value, str):
                        cell.data_type = "s"
                    cells.append(cell)
                ws.append(cells)
            workbook.save(temp)
            workbook.close()
            notes.append("Text cells are explicitly stored as text, including values beginning with '='.")
        if temp.stat().st_size > MAX_FILE_BYTES:
            raise ValueError("The exported file exceeds the 64 MiB limit.")
        os.replace(temp, path)
        return {"path": str(path), "rows": len(data["rows"]), "columns": len(data["columns"]), "notes": notes}
    finally:
        if temporary:
            Path(temporary).unlink(missing_ok=True)


def export_spjs(path, workbook):
    """Atomically export a complete workbook to the legacy SPJS format."""
    from backend.spjs_import import serialize_workbook
    if not isinstance(path, str) or not path.strip():
        raise ValueError("An output path is required.")
    output = Path(path).expanduser().resolve()
    if output.suffix.lower() != ".spjs" or not output.parent.is_dir():
        raise ValueError("Choose an existing output folder and an .spjs filename.")
    document = serialize_workbook(workbook)
    temporary = None
    try:
        fd, temporary = tempfile.mkstemp(prefix=".statworkbench-", suffix=".spjs", dir=output.parent)
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            json.dump(document, stream, ensure_ascii=False, allow_nan=False)
            stream.write("\n")
        if Path(temporary).stat().st_size > MAX_FILE_BYTES:
            raise ValueError("The exported file exceeds the 64 MiB limit.")
        os.replace(temporary, output)
        return {"path": str(output), "sheets": len(document["workbook"]["sheets"]), "notes": ["SPJS contains current worksheets and supported chart data, styles and placement. Save Project also preserves analysis history and unsupported drawings.", *document.get("statworkbenchExportWarnings", [])]}
    finally:
        if temporary:
            Path(temporary).unlink(missing_ok=True)
