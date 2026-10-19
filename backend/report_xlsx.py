"""Bounded Excel report export with self-contained editable chart series."""
from __future__ import annotations

import base64
import binascii
import math
import os
from pathlib import Path
import re
import struct
import tempfile
import zlib

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, ScatterChart
from openpyxl.chart.data_source import AxDataSource, NumData, NumDataSource, NumVal, StrData, StrVal
from openpyxl.chart.marker import Marker
from openpyxl.chart.label import DataLabel, DataLabelList
from openpyxl.chart.legend import LegendEntry
from openpyxl.chart.series import DataPoint, Series, SeriesLabel, XYSeries
from openpyxl.chart.text import RichText, Text
from openpyxl.descriptors import Typed
from openpyxl.drawing.image import Image
from openpyxl.drawing.text import CharacterProperties, Paragraph, RegularTextRun
from openpyxl.styles import Font
from openpyxl.utils.cell import range_boundaries

MAX_BYTES = 64 * 1024 * 1024
MAX_POINTS = 200_000
PALETTE = ["4869df", "28a494", "e3a64f", "b274ca", "d96d80", "579bc1"]
KINDS = {"line", "spline", "scatter", "step", "bar", "histogram", "box"}
CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def _text(value, limit=200):
    if not isinstance(value, str) or len(value) > limit or CONTROL.search(value):
        raise ValueError("Invalid report text.")
    return value


def _num(values):
    return NumDataSource(numLit=NumData(ptCount=len(values), pt=[NumVal(idx=i, v=value) for i, value in enumerate(values) if _number(value)]))


def _categories(values):
    if all(_number(value) for value in values):
        return AxDataSource(numLit=_num(values).numLit)
    return AxDataSource(strLit=StrData(ptCount=len(values), pt=[StrVal(idx=i, v=str(value)) for i, value in enumerate(values)]))


class _TextLabel(DataLabel):
    # OOXML CT_DLbl supports literal rich text; openpyxl's public model omits tx.
    tx = Typed(expected_type=Text, allow_none=True)
    __elements__ = ("idx", "tx") + DataLabel.__elements__[1:]
    __nested__ = DataLabel.__nested__

    def __init__(self, idx, text):
        super().__init__(idx=idx, dLblPos="b", showVal=False, showSerName=False, showLegendKey=False)
        self.tx = Text(rich=RichText(p=[Paragraph(r=[RegularTextRun(rPr=CharacterProperties(sz=900, solidFill="8A99AF"), t=text)])]))


class _PNG(Image):
    """Raw validated PNG embedding uses no Pillow decoding or dependency."""
    def __init__(self, data_url):
        if not isinstance(data_url, str) or len(data_url) > 12 * 1024 * 1024 or not data_url.startswith("data:image/png;base64,"):
            raise ValueError("A PNG fallback is required for this chart.")
        try:
            payload = base64.b64decode(data_url.split(",", 1)[1], validate=True)
        except (ValueError, binascii.Error) as exc:
            raise ValueError("Invalid PNG chart encoding.") from exc
        if len(payload) > 8 * 1024 * 1024 or payload[:8] != b"\x89PNG\r\n\x1a\n":
            raise ValueError("Invalid PNG chart.")
        offset, dimensions, image_data, finished = 8, None, False, False
        while offset + 12 <= len(payload):
            length = struct.unpack(">I", payload[offset:offset + 4])[0]
            kind = payload[offset + 4:offset + 8]
            end = offset + 12 + length
            if end > len(payload):
                raise ValueError("Truncated PNG chart.")
            body = payload[offset + 8:offset + 8 + length]
            if zlib.crc32(kind + body) & 0xffffffff != struct.unpack(">I", payload[end - 4:end])[0]:
                raise ValueError("Invalid PNG chart checksum.")
            if dimensions is None:
                if kind != b"IHDR" or length != 13:
                    raise ValueError("Missing PNG chart dimensions.")
                width, height, depth, color, compression, filtering, interlace = struct.unpack(">IIBBBBB", body)
                depths = {0: {1, 2, 4, 8, 16}, 2: {8, 16}, 3: {1, 2, 4, 8}, 4: {8, 16}, 6: {8, 16}}
                if not 0 < width <= 8192 or not 0 < height <= 8192 or width * height > 16_777_216 or depth not in depths.get(color, set()) or compression or filtering or interlace not in {0, 1}:
                    raise ValueError("PNG chart dimensions or encoding exceed supported limits.")
                dimensions = (width, height)
            elif kind == b"IHDR":
                raise ValueError("Duplicate PNG chart header.")
            if kind == b"IDAT":
                image_data = True
            if kind == b"IEND":
                if length or end != len(payload):
                    raise ValueError("Invalid PNG chart ending.")
                finished = True
                break
            if kind not in {b"IHDR", b"PLTE", b"IDAT", b"IEND"} and not (kind[0] & 32):
                raise ValueError("Unsupported PNG critical chunk.")
            offset = end
        if not dimensions or not image_data or not finished:
            raise ValueError("Incomplete PNG chart.")
        self.width, self.height = dimensions
        self.format, self.ref, self.payload = "png", None, payload

    def _data(self):
        return self.payload


def _validated_chart(chart):
    if not isinstance(chart, dict) or chart.get("type") not in KINDS:
        raise ValueError("Invalid report chart type.")
    for key in ("title", "xLabel", "yLabel"):
        if key in chart:
            _text(chart[key])
    if "xTicks" in chart:
        ticks = chart["xTicks"]
        if not isinstance(ticks, list) or len(ticks) > 500:
            raise ValueError("Invalid chart tick labels.")
        for tick in ticks:
            if not isinstance(tick, dict) or not _number(tick.get("value")):
                raise ValueError("Invalid chart tick coordinate.")
            _text(tick.get("label"), 1000)
    for axis in ("x", "y"):
        lower, upper = chart.get(axis + "Min"), chart.get(axis + "Max")
        if lower is not None or upper is not None:
            if not _number(lower) or not _number(upper) or upper <= lower:
                raise ValueError("Invalid chart axis limits.")
        major = chart.get(axis + "MajorUnit")
        if major is not None and (not _number(major) or major <= 0):
            raise ValueError("Chart axis major units must be positive finite numbers.")
    if not isinstance(chart.get("series"), list) or not 0 < len(chart["series"]) <= 100:
        raise ValueError("A chart requires 1 to 100 series.")
    series, total = [], 0
    for index, source in enumerate(chart["series"]):
        if not isinstance(source, dict) or source.get("type", chart["type"]) not in KINDS:
            raise ValueError("Invalid chart series type.")
        ys = source.get("y")
        if not isinstance(ys, list) or not ys or any(y is not None and not _number(y) for y in ys):
            raise ValueError("Chart values must be finite numbers or missing values.")
        xs = source.get("x", list(range(1, len(ys) + 1)))
        if not isinstance(xs, list) or len(xs) != len(ys) or any(x is not None and not _number(x) and not isinstance(x, str) for x in xs):
            raise ValueError("Chart X and Y dimensions must agree.")
        for x in xs:
            if isinstance(x, str):
                _text(x, 1000)
        total += len(ys)
        if total > MAX_POINTS:
            raise ValueError("Chart exceeds the 200,000 point limit.")
        kind = source.get("type", chart["type"])
        edges = source.get("binEdges", chart.get("binEdges"))
        if kind == "histogram" and (not isinstance(edges, list) or len(edges) != len(ys) + 1 or any(not _number(v) for v in edges) or any(b <= a for a, b in zip(edges, edges[1:])) or any(not _number(x) or not edges[i] <= x <= edges[i + 1] for i, x in enumerate(xs))):
            raise ValueError("Histogram bins must have finite, increasing boundaries.")
        color = source.get("color", "#" + PALETTE[index % len(PALETTE)])
        if not isinstance(color, str) or not re.fullmatch(r"#[0-9a-fA-F]{6}", color):
            raise ValueError("Invalid chart color.")
        markers = source.get("markers", {})
        if not isinstance(markers, dict):
            raise ValueError("Invalid chart markers.")
        for key, value, minimum, maximum in (("width", source.get("width", 1.7), .25, 12), ("marker radius", markers.get("radius", 3), 1, 20)):
            if not _number(value) or not minimum <= value <= maximum:
                raise ValueError("Invalid chart " + key + ".")
        series.append({**source, "type": kind, "x": xs, "y": ys, "color": color[1:], "name": _text(source.get("name", "Series " + str(index + 1))), "markers": markers, "binEdges": edges})
    return series, total


def _style(series, source):
    color = source["color"]
    series.graphicalProperties.line.solidFill = color
    series.graphicalProperties.line.width = round(source.get("width", 1.7) * 9525)
    series.graphicalProperties.line.prstDash = {"dash": "dash", "dot": "sysDot", "dashdot": "dashDot"}.get(source.get("dash"), "solid")
    series.graphicalProperties.solidFill = color
    if source["type"] == "scatter":
        series.graphicalProperties.line.solidFill = None
        series.graphicalProperties.line.noFill = True
    markers = source["markers"]
    visible = markers.get("visible", source["type"] == "scatter" or len(source["y"]) < 45)
    series.marker = Marker(symbol="circle" if visible else "none", size=max(2, round(markers.get("radius", 3 if source["type"] == "scatter" else 2.2) * 1.5)))
    series.marker.graphicalProperties.line.solidFill = color
    series.marker.graphicalProperties.solidFill = "FFFFFF" if markers.get("style") == "nofill" else color
    series.smooth = source["type"] == "spline"


def _axis_style(chart, spec, numeric_x):
    chart.title = spec.get("title", "Chart")
    chart.title.overlay = False
    chart.x_axis.title = spec.get("xLabel")
    chart.y_axis.title = spec.get("yLabel")
    chart.legend = chart.legend if spec.get("legend", True) else None
    if chart.legend:
        chart.legend.overlay = False
    chart.x_axis.axPos, chart.y_axis.axPos = "b", "l"
    chart.x_axis.tickLblPos = chart.y_axis.tickLblPos = "nextTo"
    for axis in (chart.x_axis, chart.y_axis):
        axis.delete = False
        axis.numFmt = "General"
        axis.crosses = "min"
    chart.visible_cells_only = False
    chart.display_blanks = "gap"
    chart.width, chart.height = 18, 11.1
    for name, axis in (("y", chart.y_axis), ("x", chart.x_axis)):
        if name == "x" and not numeric_x:
            axis.scaling.orientation = "maxMin" if spec.get("reverseX") else "minMax"
            continue
        if name + "Min" in spec:
            axis.scaling.min, axis.scaling.max = spec[name + "Min"], spec[name + "Max"]
        if spec.get(name + "MajorUnit") is not None:
            axis.majorUnit = spec[name + "MajorUnit"]
        axis.scaling.orientation = "maxMin" if spec.get("reverse" + name.upper()) else "minMax"


def _numeric_bounds(spec, sources):
    xs = [x for source in sources for x in source["x"] if _number(x)]
    ys = [y for source in sources for y in source["y"] if _number(y)]
    if not xs or not ys:
        raise ValueError("A chart requires finite coordinates.")
    xmin, xmax, ymin, ymax = min(xs), max(xs), min(ys), max(ys)
    if xmin == xmax:
        xmin -= .5; xmax += .5
    if ymin == ymax:
        ymin -= 1; ymax += 1
    padding = (ymax - ymin) * .08
    return (spec.get("xMin", xmin), spec.get("xMax", xmax), spec.get("yMin", ymin - padding), spec.get("yMax", ymax + padding))


def _tick_labels(chart, spec, sources):
    if not spec.get("xTicks"):
        return
    xmin, xmax, ymin, ymax = _numeric_bounds(spec, sources)
    ticks = [tick for tick in spec["xTicks"] if xmin <= tick["value"] <= xmax]
    if not ticks:
        return
    chart.x_axis.scaling.min, chart.x_axis.scaling.max = xmin, xmax
    chart.y_axis.scaling.min, chart.y_axis.scaling.max = ymin, ymax
    # openpyxl's NestedNoneSet maps the legal OOXML token "none" to omission;
    # Excel then restores numeric labels over our literal labels. Preserve it.
    chart.x_axis.__dict__["tickLblPos"] = "none"
    index = len(chart.series)
    series = XYSeries(idx=index, order=index, tx=SeriesLabel(v="Axis labels"), xVal=AxDataSource(numLit=_num([tick["value"] for tick in ticks]).numLit), yVal=_num([ymax if spec.get("reverseY") else ymin] * len(ticks)))
    series.graphicalProperties.line.noFill = True
    series.marker = Marker(symbol="none")
    series.dLbls = DataLabelList(dLbl=[_TextLabel(idx=i, text=tick["label"]) for i, tick in enumerate(ticks)])
    chart.series.append(series)
    if chart.legend:
        chart.legend.legendEntry.append(LegendEntry(idx=index, delete=True))


def _box_chart(spec, sources):
    """Represent box plots with ordinary editable XY geometry."""
    if any(source["type"] != "box" for source in sources):
        return None, ["Mixed box/chart geometry retains its image."]
    outlines, ticks, extent_values = [], [], []
    width = min(len(sources) / 12, .45)  # Match the renderer's 50 px cap in its 600 px plot.
    for index, source in enumerate(sources):
        values = sorted(value for value in source["y"] if _number(value))
        if not values:
            continue
        extent_values.extend(values)
        box = source.get("box", {})
        if not isinstance(box, dict):
            raise ValueError("Invalid box plot statistics.")
        for field in ("q1", "q3", "median", "mean", "low", "high"):
            if field in box and not _number(box[field]):
                raise ValueError("Box plot statistics must be finite numbers.")
        for field in ("outliers", "mildOutliers", "extremeOutliers"):
            if field in box and (not isinstance(box[field], list) or len(box[field]) > MAX_POINTS or any(not _number(value) for value in box[field])):
                raise ValueError("Invalid box plot outliers.")
        def quantile(p):
            position = (len(values) - 1) * p
            lower = math.floor(position)
            return values[lower] + (values[min(lower + 1, len(values) - 1)] - values[lower]) * (position - lower)
        q1, median, q3 = (box.get(key, quantile(p)) for key, p in (("q1", .25), ("median", .5), ("q3", .75)))
        lo = box.get("low", next((v for v in values if v >= q1 - 1.5 * (q3 - q1)), values[0]))
        hi = box.get("high", next((v for v in reversed(values) if v <= q3 + 1.5 * (q3 - q1)), values[-1]))
        extent_values.extend([lo, hi, q1, median, q3])
        if "mean" in box:
            extent_values.append(box["mean"])
        for field in ("outliers", "mildOutliers", "extremeOutliers"):
            extent_values.extend(box.get(field, []))
        x = index + .5
        base = {**source, "type": "line", "markers": {"visible": False}, "layer": 0}
        outlines.append({**base, "x": [x, x, None, x-width/3, x+width/3, None, x-width/3, x+width/3, None, x-width/2, x-width/2, x+width/2, x+width/2, x-width/2], "y": [lo, hi, None, lo, lo, None, hi, hi, None, q1, q3, q3, q1, q1]})
        for field, value, color in (("Median", median, "2563EB"), ("Mean", box.get("mean"), "DC2626")):
            if value is not None:
                outlines.append({**base, "name": source["name"][:180] + " · " + field, "x": [x-width/2, x+width/2], "y": [value, value], "color": color, "width": 2, "_legendHidden": True})
        mild = box.get("mildOutliers", box.get("outliers", [v for v in values if v < lo or v > hi]))[:1000]
        extreme = box.get("extremeOutliers", [])[:1000]
        for kind, points, marker in (("Mild outliers", mild, "circle"), ("Extreme outliers", extreme, "star")):
            if points:
                outlines.append({**base, "name": source["name"][:175] + " · " + kind, "type": "scatter", "x": [x] * len(points), "y": points, "markers": {"visible": True, "style": "nofill", "radius": 3}, "_markerSymbol": marker, "_legendHidden": True})
        ticks.append({"value": x, "label": source["name"]})
    if not outlines:
        raise ValueError("A box plot requires finite values.")
    if len(outlines) >= 100:
        return None, ["This box plot has too many editable component series and retains its image."]
    numeric_spec = {**spec, "type": "line", "xMin": 0, "xMax": len(sources), "xTicks": ticks}
    _, _, ymin, ymax = _numeric_bounds(spec, [{"x": [0, len(sources)], "y": extent_values}])
    numeric_spec.update(yMin=ymin, yMax=ymax)
    chart, warnings = _native_chart(numeric_spec, outlines)
    return chart, warnings + ["Box plots use editable outlines, whiskers, median, mean and outlier markers; translucent box fill is omitted."]


def _native_chart(spec, sources):
    """Return (chart, warnings); unsupported geometry is deliberately a PNG."""
    if spec["type"] == "box" or any(s["type"] == "box" for s in sources):
        return _box_chart(spec, sources)
    sources = sorted(sources, key=lambda source: source.get("layer") if _number(source.get("layer")) else 0)
    bars = [s for s in sources if s["type"] in {"bar", "histogram"}]
    numeric = all(x is None or _number(x) for s in sources for x in s["x"])
    warnings = []
    if any(s["type"] == "spline" for s in sources):
        warnings.append("The spreadsheet uses its own editable spline interpolation, which may differ between control points.")
    if not bars and numeric:
        if spec.get("xTicks") and len(sources) >= 100:
            return None, ["This chart has too many series for editable custom labels and retains its image."]
        chart = ScatterChart(scatterStyle="lineMarker", varyColors=False)
        for index, source in enumerate(sources):
            xs, ys, inserted = [], [], []
            previous = None
            for x, y in zip(source["x"], source["y"]):
                if x is None or y is None:
                    xs.append(None); ys.append(None); previous = None
                    continue
                if source["type"] == "step" and previous is not None:
                    inserted.append(len(xs))
                    xs.append(x); ys.append(previous)
                xs.append(x); ys.append(y); previous = y
            series = XYSeries(idx=index, order=index, tx=SeriesLabel(v=source["name"]), xVal=AxDataSource(numLit=_num(xs).numLit), yVal=_num(ys))
            _style(series, source)
            if source.get("_markerSymbol"):
                series.marker.symbol = source["_markerSymbol"]
            series.dPt = [DataPoint(idx=i, marker=Marker(symbol="none")) for i in inserted]
            chart.series.append(series)
        _axis_style(chart, spec, True)
        if chart.legend:
            chart.legend.legendEntry.extend(LegendEntry(idx=i, delete=True) for i, source in enumerate(sources) if source.get("_legendHidden"))
        _tick_labels(chart, spec, sources)
        return chart, warnings
    # Category plots can only preserve category positions when all overlays
    # share the bar/category grid. Otherwise keep the accurately rendered image.
    categories = sources[0]["x"]
    if any(s["x"] != categories for s in sources) or any(x is None for x in categories) or len(set(categories)) != len(categories):
        return None, ["Different or repeated category coordinates retain their image to preserve chart geometry."]
    if any(s["type"] == "step" for s in sources):
        return None, ["Categorical step curves retain their image to preserve step geometry."]
    if "xMin" in spec or "xMax" in spec or spec.get("xTicks") or spec.get("xMajorUnit") is not None:
        return None, ["Custom category-axis coordinates retain their image to preserve tick placement."]
    if numeric and bars and len(categories) > 1:
        gaps = [b - a for a, b in zip(categories, categories[1:])]
        if gaps[0] <= 0 or any(not math.isclose(gap, gaps[0], rel_tol=1e-9, abs_tol=1e-12) for gap in gaps):
            return None, ["Unequally spaced numeric bars retain their image to preserve horizontal distances."]
    for source in bars:
        if source["type"] == "histogram":
            edges = source["binEdges"]
            widths = [b - a for a, b in zip(edges, edges[1:])]
            if any(not math.isclose(width, widths[0], rel_tol=1e-9, abs_tol=1e-12) for width in widths) or any(not math.isclose(x, (edges[i] + edges[i + 1]) / 2, rel_tol=1e-9, abs_tol=1e-12) for i, x in enumerate(categories)):
                return None, ["Unequal histogram bins retain their image to preserve bin widths and positions."]
    if bars and any(s["type"] == "histogram" for s in bars) and len(bars) != 1:
        return None, ["Overlapping histograms retain their image to preserve bin placement."]
    chart = BarChart(varyColors=False) if bars else LineChart(varyColors=False)
    if bars:
        chart.type = "col"
        chart.grouping = "stacked" if spec.get("stack") in {"normal", "percent"} else "clustered"
        chart.overlap = 100 if chart.grouping == "stacked" else 0
        chart.gapWidth = 2 if bars[0]["type"] == "histogram" else 45
    line = LineChart(varyColors=False) if bars and len(bars) != len(sources) else chart
    if bars and spec.get("stack") and line is not chart:
        return None, ["Stacked mixed charts retain their image to preserve series placement."]
    for index, source in enumerate(sources):
        ys = source["y"]
        if bars and spec.get("stack") == "percent":
            ys = [None if y is None else (100 * y / total if total else 0) for j, y in enumerate(ys) for total in [sum(abs(s["y"][j]) for s in sources if s["y"][j] is not None)]]
        series = Series(idx=index, order=index, tx=SeriesLabel(v=source["name"]), cat=_categories(categories), val=_num(ys))
        _style(series, source)
        (chart if source in bars else line).series.append(series)
    _axis_style(chart, spec, False)
    if line is not chart:
        _axis_style(line, spec, False)
        chart += line
    return chart, warnings


def export_report_xlsx(path, report):
    """Create one worksheet with inline chart data and no workbook links."""
    if not isinstance(path, str) or not path.strip():
        raise ValueError("An output path is required.")
    output = Path(path).expanduser().resolve()
    if output.suffix.lower() != ".xlsx" or not output.parent.is_dir():
        raise ValueError("Choose an existing output folder and an .xlsx filename.")
    if not isinstance(report, dict) or not isinstance(report.get("matrix"), list) or not 0 < len(report["matrix"]) <= 100_000:
        raise ValueError("Invalid report matrix.")
    matrix = report["matrix"]
    if any(not isinstance(row, list) for row in matrix):
        raise ValueError("Invalid report row.")
    width = max(map(len, matrix))
    if not 0 < width <= 1000 or width * len(matrix) > 2_000_000:
        raise ValueError("Report exceeds worksheet limits.")
    font = _text(report.get("font", "Arial"), 100)
    font_size = report.get("fontSize", 11)
    if not _number(font_size) or not 6 <= font_size <= 72:
        raise ValueError("Invalid report font size.")
    number_format = _text(report.get("numberFormat", "0.00000"), 100)
    images = report.get("chartImages", [])
    formats = report.get("formats", [])
    if not isinstance(images, list) or len(images) > 30 or not isinstance(formats, list) or len(formats) > 100_000:
        raise ValueError("Too many report charts or formats.")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"
    sheet.sheet_format.defaultColWidth = 18
    sheet.sheet_format.defaultRowHeight = max(15, font_size * 1.4)
    native, raster, total_points, warnings = 0, 0, 0, []
    ordinary_font, title_font = Font(name=font, size=font_size), Font(name=font, size=font_size, bold=True)
    temporary = None
    try:
        for r, row in enumerate(matrix, 1):
            for c, value in enumerate(row, 1):
                if value is not None and not isinstance(value, (str, bool)) and not _number(value):
                    raise ValueError("Report cells must be finite scalar values.")
                if isinstance(value, str):
                    _text(value, 32767)
                cell = sheet.cell(r, c, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                cell.font = title_font if r == 1 else ordinary_font
                cell.number_format = number_format
        formatted_cells = 0
        for item in formats:
            if not isinstance(item, dict) or not isinstance(item.get("range"), str) or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,5}(?::[A-Z]{1,3}[1-9][0-9]{0,5})?", item["range"]):
                raise ValueError("Invalid report format range.")
            left, top, right, bottom = range_boundaries(item["range"])
            if left > right or top > bottom or right > width or bottom > len(matrix):
                raise ValueError("Report formats must remain inside its data.")
            formatted_cells += (right - left + 1) * (bottom - top + 1)
            if formatted_cells > 2_000_000:
                raise ValueError("Report formatting exceeds the cell-operation limit.")
            fmt = _text(item.get("numberFormat"), 100)
            for row in sheet.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right):
                for cell in row:
                    cell.number_format = fmt
        for index, item in enumerate(images):
            if not isinstance(item, dict):
                raise ValueError("Invalid report chart.")
            spec = item.get("chart")
            if spec is None:
                chart, notes = None, ["This chart has no editable series specification and retains its image."]
            else:
                sources, points = _validated_chart(spec)
                total_points += points
                if total_points > MAX_POINTS:
                    raise ValueError("Report charts exceed the 200,000 point limit.")
                if report.get("app") == "numbers" and (spec.get("xTicks") or any(source["type"] == "box" for source in sources)):
                    chart, notes = None, ["Numbers does not preserve custom point labels on imported charts; this chart retains its image so its group/axis labels remain intact."]
                else:
                    chart, notes = _native_chart(spec, sources)
            label = _text(item.get("title", (spec or {}).get("title", "Chart")))
            warnings.extend(label + ": " + note for note in notes)
            anchor = "A" + str(len(matrix) + 3 + index * 25)
            if chart is None:
                image = _PNG(item.get("png"))
                image.width, image.height = 680, 420
                sheet.add_image(image, anchor)
                raster += 1
            else:
                sheet.add_chart(chart, anchor)
                native += 1
        descriptor, temporary = tempfile.mkstemp(prefix=".statworkbench-report-", suffix=".xlsx", dir=output.parent)
        os.close(descriptor)
        workbook.save(temporary)
        if Path(temporary).stat().st_size > MAX_BYTES:
            raise ValueError("The exported report exceeds the 64 MiB limit.")
        os.replace(temporary, output)
        return {"path": str(output), "nativeCharts": native, "imageCharts": raster, "rows": len(matrix), "columns": width, "warnings": warnings}
    finally:
        workbook.close()
        if temporary:
            Path(temporary).unlink(missing_ok=True)
