"""Read actual OOXML series and drawing relationships, not chart object counts."""
import base64
import json
from pathlib import Path
import struct
import tempfile
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
import zipfile
import zlib

from openpyxl import load_workbook
from backend.report_xlsx import export_report_xlsx

C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def png():
    def chunk(kind, body):
        return struct.pack(">I", len(body)) + kind + body + struct.pack(">I", zlib.crc32(kind + body) & 0xffffffff)
    data = b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 1, 8, 6, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(b"\0" + b"\x10\x20\x30\xff" * 2)) + chunk(b"IEND", b"")
    return "data:image/png;base64," + base64.b64encode(data).decode()


def values(series, branch):
    return {int(point.attrib["idx"]): float(point.find(C + "v").text) for point in series.findall(C + branch + "/" + C + "numLit/" + C + "pt")}


class ReportXlsxTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.path = Path(self.temp.name) / "report.xlsx"

    def tearDown(self):
        self.temp.cleanup()

    def export(self, charts, **options):
        report = {"matrix": [["Report", "=1+1", "00123"], [1, 2.5, True]], "font": "Arial", "fontSize": 12, "numberFormat": "0.000", "formats": [{"range": "B2", "numberFormat": "0.00%"}], "chartImages": [{"chart": chart, "png": png()} for chart in charts], **options}
        result = export_report_xlsx(str(self.path), report)
        with zipfile.ZipFile(self.path) as archive:
            xml = [ET.fromstring(archive.read(name)) for name in sorted(archive.namelist()) if name.startswith("xl/charts/chart") and name.endswith(".xml")]
            for chart in xml:
                for parent in chart.iter():
                    fills = [node for node in parent if node.tag in {A + "noFill", A + "solidFill", A + "gradFill", A + "blipFill", A + "pattFill", A + "grpFill"}]
                    self.assertLessEqual(len(fills), 1, "DrawingML fill is an exclusive choice, including line fills.")
                    if parent.tag in {C + "scatterChart", C + "barChart", C + "lineChart"}:
                        self.assertEqual(parent.find(C + "varyColors").attrib["val"], "0", "Excel must identify series rather than individual points in its legend.")
            self.assertFalse(any("externalLink" in name for name in archive.namelist()))
            for name in archive.namelist():
                if name.endswith(".xml") or name.endswith(".rels"):
                    tree = ET.fromstring(archive.read(name))
                    self.assertFalse(any(node.tag.rsplit("}", 1)[-1] == "f" for node in tree.iter()), name)
                    self.assertFalse(any(node.attrib.get("TargetMode") == "External" for node in tree.iter()), name)
        return result, xml

    def test_numeric_lines_scatter_spline_step_values_gaps_and_axis_styles(self):
        chart = {"type": "line", "title": "Dependence", "xLabel": "X input", "yLabel": "Y result", "xMin": -2, "xMax": 7, "xMajorUnit": .5, "yMin": -5, "yMax": 10, "yMajorUnit": 2, "reverseX": True, "legend": False, "series": [
            {"name": "Line", "x": [0, 2, 3, 4], "y": [1, None, 3, 2], "color": "#123456", "width": 3, "dash": "dashdot", "markers": {"visible": True, "radius": 4, "style": "nofill"}},
            {"name": "Points", "type": "scatter", "x": [0, 2], "y": [5, 8]},
            {"name": "Curve", "type": "spline", "x": [0, 2], "y": [3, 5]},
            {"name": "Survival", "type": "step", "x": [0, 2, 5], "y": [1, .8, .2]},
        ]}
        result, charts = self.export([chart])
        self.assertEqual((result["nativeCharts"], result["imageCharts"]), (1, 0))
        root = charts[0]
        series = root.findall(".//" + C + "scatterChart/" + C + "ser")
        self.assertEqual([s.find(C + "tx/" + C + "v").text for s in series], ["Line", "Points", "Curve", "Survival"])
        self.assertEqual(values(series[0], "xVal"), {0: 0, 2: 3, 3: 4})
        self.assertEqual(values(series[0], "yVal"), {0: 1, 2: 3, 3: 2})
        self.assertEqual(values(series[3], "xVal"), {0: 0, 1: 2, 2: 2, 3: 5, 4: 5})
        self.assertEqual(values(series[3], "yVal"), {0: 1, 1: 1, 2: .8, 3: .8, 4: .2})
        self.assertIsNotNone(series[1].find(C + "spPr/" + A + "ln/" + A + "noFill"))
        self.assertEqual(series[2].find(C + "smooth").attrib["val"], "1")
        self.assertEqual(series[0].find(C + "spPr/" + A + "ln").attrib["w"], "28575")
        self.assertEqual(series[0].find(C + "spPr/" + A + "ln/" + A + "prstDash").attrib["val"], "dashDot")
        self.assertEqual(series[0].find(C + "marker/" + C + "size").attrib["val"], "6")
        self.assertEqual(series[0].find(C + "marker/" + C + "spPr/" + A + "solidFill/" + A + "srgbClr").attrib["val"], "FFFFFF")
        axes = root.findall(".//" + C + "valAx")
        self.assertEqual([(a.find(C + "scaling/" + C + "min").attrib["val"], a.find(C + "scaling/" + C + "max").attrib["val"]) for a in axes], [("-2", "7"), ("-5", "10")])
        self.assertEqual(axes[0].find(C + "scaling/" + C + "orientation").attrib["val"], "maxMin")
        self.assertEqual([a.find(C + "majorUnit").attrib["val"] for a in axes], ["0.5", "2"])
        self.assertIsNone(root.find(".//" + C + "legend"))
        self.assertEqual(root.find(".//" + C + "dispBlanksAs").attrib["val"], "gap")

    def test_category_bar_line_and_percent_stacks_have_literal_series(self):
        charts = [
            {"type": "bar", "title": "Groups", "series": [{"name": "Count", "x": ["A", "B"], "y": [2, 4]}, {"name": "Fit", "type": "line", "x": ["A", "B"], "y": [1, 3]}]},
            {"type": "bar", "stack": "percent", "series": [{"x": ["A"], "y": [-1]}, {"x": ["A"], "y": [3]}]},
        ]
        _, roots = self.export(charts)
        bar = roots[0].find(".//" + C + "barChart/" + C + "ser")
        line = roots[0].find(".//" + C + "lineChart/" + C + "ser")
        self.assertEqual([v.text for v in bar.findall(C + "cat/" + C + "strLit/" + C + "pt/" + C + "v")], ["A", "B"])
        self.assertEqual(values(bar, "val"), {0: 2, 1: 4})
        self.assertEqual(values(line, "val"), {0: 1, 1: 3})
        stacked = roots[1].find(".//" + C + "barChart")
        self.assertEqual(stacked.find(C + "grouping").attrib["val"], "stacked")
        self.assertEqual([values(s, "val")[0] for s in stacked.findall(C + "ser")], [-25, 75])

    def test_histogram_bins_are_native_when_geometry_is_representable(self):
        result, roots = self.export([{"type": "histogram", "yMin": 0, "yMax": 5, "series": [{"x": [.5, 1.5], "y": [2, 3], "binEdges": [0, 1, 2]}]}])
        self.assertEqual(result["nativeCharts"], 1)
        plot = roots[0].find(".//" + C + "barChart")
        self.assertEqual(plot.find(C + "gapWidth").attrib["val"], "2")
        self.assertEqual(values(plot.find(C + "ser"), "val"), {0: 2, 1: 3})
        self.assertEqual(values(plot.find(C + "ser"), "cat"), {0: .5, 1: 1.5})

    def test_unsupported_geometry_embeds_exact_png_with_explicit_warnings(self):
        report = [{"type": "line", "series": [{"x": ["A", "A"], "y": [1, 2]}]}, {"type": "histogram", "series": [{"x": [.5, 2], "y": [2, 3], "binEdges": [0, 1, 3]}]}]
        result, roots = self.export(report)
        self.assertEqual((result["nativeCharts"], result["imageCharts"]), (0, 2))
        self.assertEqual(len(result["warnings"]), 2)
        self.assertEqual(roots, [])
        with zipfile.ZipFile(self.path) as archive:
            images = [name for name in archive.namelist() if name.startswith("xl/media/")]
            self.assertEqual(len(images), 2)
            self.assertTrue(all(archive.read(name) == base64.b64decode(png().split(",")[1]) for name in images))
            rels = archive.read("xl/drawings/_rels/drawing1.xml.rels").decode()
            self.assertIn("/image", rels)
            self.assertNotIn("/chart", rels)

    def test_descending_numeric_bars_and_category_major_units_preserve_image_geometry(self):
        result, roots = self.export([
            {"type": "bar", "series": [{"x": [2, 1], "y": [3, 4]}]},
            {"type": "bar", "xMajorUnit": .5, "series": [{"x": [1, 2], "y": [3, 4]}]},
        ])
        self.assertEqual((result["nativeCharts"], result["imageCharts"]), (0, 2))
        self.assertEqual(roots, [])
        self.assertIn("horizontal distances", result["warnings"][0])
        self.assertIn("tick placement", result["warnings"][1])

    def test_numeric_custom_ticks_are_literal_labels_and_reference_layer_precedes_points(self):
        sources = [{"name": "Observed", "type": "scatter", "x": [1, 2], "y": [3, 4], "markers": {"style": "nofill"}}, {"name": "Reference", "type": "line", "x": [1, 2], "y": [2, 5], "layer": -1}]
        result, roots = self.export([
            {"type": "scatter", "xTicks": [{"value": 1, "label": "First"}], "series": sources},
            {"type": "scatter", "series": sources},
        ])
        self.assertEqual((result["nativeCharts"], result["imageCharts"]), (2, 0))
        labels = roots[0].findall(".//" + C + "scatterChart/" + C + "ser")[-1]
        self.assertEqual(values(labels, "xVal"), {0: 1})
        self.assertAlmostEqual(values(labels, "yVal")[0], 1.76)
        self.assertEqual(labels.find(".//" + A + "t").text, "First")
        self.assertIsNotNone(labels.find(C + "spPr/" + A + "ln/" + A + "noFill"))
        self.assertEqual(roots[0].find(".//" + C + "legendEntry/" + C + "delete").attrib["val"], "1")
        self.assertEqual(roots[0].find(".//" + C + "valAx/" + C + "tickLblPos").attrib["val"], "none")
        series = roots[1].findall(".//" + C + "scatterChart/" + C + "ser")
        self.assertEqual([s.find(C + "tx/" + C + "v").text for s in series], ["Reference", "Observed"])
        self.assertEqual([s.find(C + "order").attrib["val"] for s in series], ["0", "1"])
        self.assertEqual([source["name"] for source in sources], ["Observed", "Reference"])

    def test_box_plot_preserves_supplied_statistics_and_distinct_outliers_as_xy_geometry(self):
        result, roots = self.export([{"type": "box", "yMin": -2, "yMax": 20, "series": [{"name": "Group", "y": [1, 2, 7, 9], "box": {"q1": 3, "median": 5, "q3": 8, "low": 1, "high": 10, "mean": 6, "mildOutliers": [12], "extremeOutliers": [18]}}]}])
        self.assertEqual((result["nativeCharts"], result["imageCharts"]), (1, 0))
        self.assertIn("translucent box fill", result["warnings"][0])
        series = roots[0].findall(".//" + C + "scatterChart/" + C + "ser")
        self.assertEqual([s.find(C + "tx/" + C + "v").text for s in series], ["Group", "Group · Median", "Group · Mean", "Group · Mild outliers", "Group · Extreme outliers", "Axis labels"])
        self.assertEqual(values(series[0], "yVal"), {0: 1, 1: 10, 3: 1, 4: 1, 6: 10, 7: 10, 9: 3, 10: 8, 11: 8, 12: 3, 13: 3})
        self.assertEqual(values(series[1], "yVal"), {0: 5, 1: 5})
        self.assertEqual(values(series[2], "yVal"), {0: 6, 1: 6})
        self.assertEqual(values(series[3], "yVal"), {0: 12})
        self.assertEqual(values(series[4], "yVal"), {0: 18})
        self.assertEqual(series[4].find(C + "marker/" + C + "symbol").attrib["val"], "star")
        self.assertEqual(series[5].find(".//" + A + "t").text, "Group")
        self.assertEqual(values(series[5], "xVal"), {0: .5})
        self.assertEqual(values(series[5], "yVal"), {0: -2})
        axes = roots[0].findall(".//" + C + "valAx")
        self.assertEqual([axis.find(C + "axPos").attrib["val"] for axis in axes], ["b", "l"])
        self.assertEqual(axes[1].find(C + "tickLblPos").attrib["val"], "nextTo")
        self.assertEqual([axis.find(C + "delete").attrib["val"] for axis in axes], ["0", "0"])

    def test_cells_are_typed_inert_and_number_formats_survive(self):
        self.export([])
        workbook = load_workbook(self.path)
        try:
            sheet = workbook.active
            self.assertEqual((sheet["B1"].value, sheet["B1"].data_type), ("=1+1", "s"))
            self.assertEqual(sheet["C1"].value, "00123")
            self.assertEqual(sheet["C2"].data_type, "b")
            self.assertEqual(sheet["B2"].number_format, "0.00%")
            self.assertEqual((sheet["B2"].font.name, sheet["B2"].font.sz), ("Arial", 12))
            self.assertEqual(workbook.sheetnames, ["Report"])
        finally:
            workbook.close()

    def test_numbers_preserves_custom_axis_and_box_group_labels_as_images(self):
        result, roots = self.export([
            {"type": "scatter", "xTicks": [{"value": 1, "label": "First"}], "series": [{"x": [1, 2], "y": [2, 4]}]},
            {"type": "box", "series": [{"name": "Group", "y": [1, 2, 4]}]},
        ], app="numbers")
        self.assertEqual((result["nativeCharts"], result["imageCharts"]), (0, 2))
        self.assertEqual(roots, [])
        self.assertTrue(all("Numbers does not preserve custom point labels" in warning for warning in result["warnings"]))

    def test_box_default_quartiles_and_raw_data_axis_extent(self):
        _, roots = self.export([{"type": "box", "series": [{"y": [1, 2, 3, 4, 100], "box": {"outliers": []}}]}])
        series = roots[0].findall(".//" + C + "scatterChart/" + C + "ser")
        self.assertEqual(values(series[0], "yVal"), {0: 1, 1: 4, 3: 1, 4: 1, 6: 4, 7: 4, 9: 2, 10: 4, 11: 4, 12: 2, 13: 2})
        self.assertEqual(values(series[1], "yVal"), {0: 3, 1: 3})
        axis = roots[0].findall(".//" + C + "valAx")[1]
        self.assertAlmostEqual(float(axis.find(C + "scaling/" + C + "max").attrib["val"]), 107.92)

    def test_box_separate_outlier_arrays_expand_automatic_axis_bounds(self):
        _, roots = self.export([{"type": "box", "series": [{"y": [1, 2, 3], "box": {"mildOutliers": [-10], "extremeOutliers": [20]}}]}])
        axis = roots[0].findall(".//" + C + "valAx")[1]
        self.assertAlmostEqual(float(axis.find(C + "scaling/" + C + "min").attrib["val"]), -12.4)
        self.assertAlmostEqual(float(axis.find(C + "scaling/" + C + "max").attrib["val"]), 22.4)
        series = roots[0].findall(".//" + C + "scatterChart/" + C + "ser")
        self.assertEqual(values(series[2], "yVal"), {0: -10})
        self.assertEqual(values(series[3], "yVal"), {0: 20})

    def test_bad_data_crc_bounds_and_atomic_failures_are_rejected(self):
        self.path.write_bytes(b"original")
        invalid = [
            {"matrix": [[float("inf")]]},
            {"matrix": [[1]], "formats": [{"range": "B2", "numberFormat": "0"}]},
            {"matrix": [[1]], "chartImages": [{"chart": {"type": "line", "series": [{"x": [1], "y": [1, 2]}]}}]},
            {"matrix": [[1]], "chartImages": [{"chart": {"type": "line", "xMin": 1, "xMax": 0, "series": [{"y": [1]}]}}]},
            *[{"matrix": [[1]], "chartImages": [{"chart": {"type": "line", "xMajorUnit": value, "series": [{"y": [1]}]}}]} for value in (0, -1, float("inf"), True)],
            *[{"matrix": [[1]], "chartImages": [{"chart": {"type": "line", "xTicks": ticks, "series": [{"y": [1]}]}}]} for ticks in ({}, [{"value": 1, "label": 7}], [{"value": float("inf"), "label": "Bad"}])],
            {"matrix": [[1]], "chartImages": [{"chart": {"type": "box", "series": [{"y": [1], "box": {"q1": float("nan")}}]}}]},
            {"matrix": [[1]], "chartImages": [{"png": png()[:-8] + "AAAAAAAA"}]},
        ]
        for report in invalid:
            with self.subTest(report=report), self.assertRaises(ValueError):
                export_report_xlsx(str(self.path), report)
            self.assertEqual(self.path.read_bytes(), b"original")
        with patch("backend.report_xlsx.MAX_BYTES", 1), self.assertRaisesRegex(ValueError, "64 MiB"):
            export_report_xlsx(str(self.path), {"matrix": [[1]]})
        self.assertEqual(self.path.read_bytes(), b"original")
        self.assertEqual(list(self.path.parent.glob(".statworkbench-report-*")), [])


if __name__ == "__main__":
    unittest.main()
