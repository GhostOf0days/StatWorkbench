"""Independent file-boundary and round-trip checks."""

import hashlib
from contextlib import closing
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

from backend.common import validate_dataset
from backend.file_io import import_file, export_file


class FileIOTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)

    def tearDown(self):
        self.directory.cleanup()

    def write(self, name, content):
        path = self.root / name
        path.write_text(content, encoding="utf-8", newline="")
        return str(path)

    def test_csv_preserves_missing_positions_and_identifier_text(self):
        path = self.write("values.csv", "ID,Value,Label\n001,1.25,NA\n002,,NaN\n003,-3,hello\n,,\n")
        data = import_file(path)
        self.assertEqual(data["rows"], [["001", 1.25, "NA"], ["002", None, "NaN"], ["003", -3, "hello"], [None, None, None]])

    def test_csv_multiline_quoted_and_large_integer(self):
        data = import_file(self.write("quoted.csv", 'ID,Text\n9007199254740993,"line one\nline two"\n'))
        self.assertEqual(data["rows"], [["9007199254740993", "line one\nline two"]])

    def test_csv_preserves_crlf_inside_quoted_fields(self):
        data = import_file(self.write("quoted-crlf.csv", 'ID,Text\r\n1,"line one\r\nline two"\r\n'))
        self.assertEqual(data["rows"], [[1, "line one\r\nline two"]])

    def test_delimited_ragged_and_duplicate_headers_rejected(self):
        with self.assertRaisesRegex(ValueError, "Duplicate"):
            import_file(self.write("dupe.csv", "x,x\n1,2\n"))
        with self.assertRaisesRegex(ValueError, "fields"):
            import_file(self.write("ragged.tsv", "a\tb\n1\t2\t3\n"))

    def test_json_round_trip_preserves_empty_string_and_null(self):
        data = {"columns": ["x", "flag"], "rows": [["", False], [None, True], [1.5, None]], "name": "日本語"}
        path = str(self.root / "data.json")
        export_file(path, data)
        self.assertEqual(import_file(path), data)

    def test_json_nonstandard_numbers_and_duplicate_keys_rejected(self):
        with self.assertRaisesRegex(ValueError, "Nonstandard"):
            import_file(self.write("nan.json", '{"columns":["x"],"rows":[[NaN]]}'))
        with self.assertRaisesRegex(ValueError, "Duplicate JSON"):
            import_file(self.write("dupe.json", '{"columns":["x"],"rows":[],"rows":[]}'))

    def test_xlsx_keeps_formula_like_text_inert(self):
        from openpyxl import load_workbook
        path = str(self.root / "values.xlsx")
        data = {"columns": ["x", "y"], "rows": [["=1+1", 2], ["@SUM(A1)", None], ["ordinary", 3]], "name": "Workbook"}
        export_file(path, data)
        wb = load_workbook(path, data_only=False)
        self.assertEqual(wb.active["A2"].data_type, "s")
        self.assertEqual(wb.active["A2"].value, "=1+1")
        wb.close()
        self.assertEqual(import_file(path)["rows"], data["rows"])

    def test_xlsx_select_sheet_and_cached_formula_policy(self):
        from openpyxl import Workbook
        path = self.root / "sheets.xlsx"
        wb = Workbook()
        wb.active.title = "First"
        wb.active.append(["x"])
        wb.active.append([1])
        second = wb.create_sheet("Second")
        second.append(["y"])
        second.append(["=1+1"])
        wb.save(path)
        wb.close()
        data = import_file(str(path), "Second")
        self.assertEqual(data["sheets"], ["First", "Second"])
        self.assertEqual(data["rows"], [[None]])
        self.assertIn("not evaluated", data["notes"][0])

    def test_sqlite_uses_explicit_existing_tables_and_is_read_only(self):
        path = self.root / "example.sqlite"
        with closing(sqlite3.connect(path)) as connection:
            connection.execute('CREATE TABLE "a strange table" (value REAL, label TEXT)')
            connection.execute('INSERT INTO "a strange table" VALUES (?, ?)', (2.5, None))
            connection.execute('CREATE TABLE other (id INTEGER)')
            connection.commit()
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        result = import_file(str(path))
        self.assertTrue(result["requires_sheet"])
        result = import_file(str(path), "a strange table")
        self.assertEqual(result["rows"], [[2.5, None]])
        with self.assertRaisesRegex(ValueError, "existing SQLite table"):
            import_file(str(path), 'other"; DROP TABLE other; --')
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), before)

    def test_sqlite_closes_connections_after_success_selection_and_error(self):
        path = self.root / "handles.sqlite"
        with closing(sqlite3.connect(path)) as connection:
            connection.execute("CREATE TABLE first (value REAL)")
            connection.execute("CREATE TABLE second (value REAL)")
        connect = sqlite3.connect
        opened = []

        def track_connection(*args, **kwargs):
            connection = connect(*args, **kwargs)
            opened.append(connection)
            return connection

        with patch("backend.file_io.sqlite3.connect", side_effect=track_connection):
            import_file(str(path))
            import_file(str(path), "first")
            with self.assertRaisesRegex(ValueError, "existing SQLite table"):
                import_file(str(path), "missing")
        self.assertEqual(len(opened), 3)
        for connection in opened:
            with self.assertRaisesRegex(sqlite3.ProgrammingError, "closed"):
                connection.execute("SELECT 1")

    def test_spjs_reads_values_without_executing_formulas_and_retains_row_zero(self):
        source = {"workbook": {"activeSheetIndex": 0, "sheets": {"Study": {"index": 0, "data": {"dataTable": {"0": {"0": {"value": "Header"}, "1": {"value": "Group"}}, "1": {"0": {"value": 3}, "1": {"formula": "DANGEROUS()"}}, "2": {"0": {"formula": "1+1", "value": 2}, "1": {"value": "A"}}}}}}}, "charts": []}
        path = self.write("study.spjs", json.dumps(source))
        data = import_file(path)
        self.assertEqual(data["columns"], ["A", "B"])
        self.assertEqual(data["rows"], [["Header", "Group"], [3, None], [2, "A"]])
        self.assertTrue(any("never evaluated" in n for n in data["notes"]))

    def test_spjs_sparse_huge_indices_rejected_before_allocation(self):
        source = {"workbook": {"sheets": {"S": {"data": {"dataTable": {"999999999": {"0": {"value": 1}}}}}}}}
        with self.assertRaisesRegex(ValueError, "limits"):
            import_file(self.write("huge.spjs", json.dumps(source)))

    def test_export_failure_keeps_original_and_removes_temp(self):
        path = self.root / "keep.json"
        path.write_text("original", encoding="utf-8")
        with patch("backend.file_io.MAX_FILE_BYTES", 1):
            with self.assertRaisesRegex(ValueError, "64 MiB"):
                export_file(str(path), {"columns": ["x"], "rows": [[1]]})
        self.assertEqual(path.read_text(), "original")
        self.assertEqual(list(self.root.glob(".statworkbench-*")), [])

    def test_dataset_limits_nested_cells_and_nonfinite_values(self):
        for value in [{"nested": 1}, [1], float("inf")]:
            with self.assertRaises(ValueError):
                validate_dataset({"columns": ["x"], "rows": [[value]]})
        with patch("backend.common.MAX_CELLS", 1):
            with self.assertRaisesRegex(ValueError, "limit"):
                validate_dataset({"columns": ["x", "y"], "rows": [[1, 2]]})

    def test_no_project_overwrite_via_dataset_export(self):
        with self.assertRaisesRegex(ValueError, "Save Project"):
            export_file(str(self.root / "study.statworkbench.json"), {"columns": ["x"], "rows": []})


if __name__ == "__main__":
    unittest.main()
